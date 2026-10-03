"""Registro dei rischi secondo la procedura di risk management: valutazione nel
ciclo, copertura, approvazione, accettazione con autorità e pareri, piani e
verifica, rischi ereditati, export ed escalation."""
import datetime
import io

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.utils import timezone
from rest_framework.exceptions import PermissionDenied
from rest_framework.test import APIClient

from apps.risk import services
from apps.risk.models import (
    RiskAcceptance,
    RiskAssessment,
    RiskAssessmentCycle,
    RiskGovernancePolicy,
    RiskMitigationPlan,
    ThreatCatalogEntry,
)

User = get_user_model()


# ── fixture ──────────────────────────────────────────────────────────────────

def _plant(code, **extra):
    from apps.plants.models import Plant

    values = {"name": f"Plant {code}", "country": "IT", "nis2_scope": "importante", "status": "attivo"}
    values.update(extra)
    return Plant.objects.create(code=code, **values)


def _user(username, role="risk_manager", scope="org", plants=()):
    from apps.auth_grc.models import UserPlantAccess

    user = User.objects.create_user(username=username, email=f"{username}@test.com", password="x",
                                    first_name=username.title())
    access = UserPlantAccess.objects.create(user=user, role=role, scope_type=scope)
    for p in plants:
        access.scope_plants.add(p)
    return user


def _client(user):
    c = APIClient()
    c.force_authenticate(user=user)
    return c


@pytest.fixture
def plant(db):
    return _plant("RG-A")


@pytest.fixture
def other_plant(db):
    return _plant("RG-B", country="PL")


@pytest.fixture
def org_user(db):
    return _user("orgrm")


@pytest.fixture
def site_user(db, plant):
    return _user("siterm", scope="single_plant", plants=[plant])


@pytest.fixture
def threats(db):
    def mk(code, types, cia):
        return ThreatCatalogEntry.objects.create(
            code=code, asset_types=types, cia=cia, source="catalog",
            translations={"it": {"title": f"Minaccia {code}"}},
        )
    return {
        "malware": mk("IN_MAL", ["IT"], ["C", "I", "A"]),
        "fire": mk("LO_FUO", ["SEDE"], ["A"]),
        "staff": mk("PE_IPR", ["PERSONALE"], ["A"]),
    }


@pytest.fixture
def cycle(org_user, plant):
    return services.start_cycle(org_user, plant, "primo")


def _objective():
    from apps.risk.models import BusinessObjective

    return BusinessObjective.objects.get_or_create(
        code="BO-T", defaults={"name": "OEM supply continuity", "impact_dimensions": ["operational"]},
    )[0]


def _eval(owner, **extra):
    data = {
        "business_objectives": [_objective()],
        "asset_type": "IT", "probability": 4, "probability_rationale": "Phishing frequente",
        "impact_operational": 4, "impact_rationale": "Fermo produzione oltre RTO",
        "owner": owner, "treatment": "mitigare", "expected_probability": 2, "expected_impact": 4,
    }
    data.update(extra)
    return data


def _completed_risk(user, plant, threat, owner=None, **extra):
    risk = services.create_risk(user, plant, {**_eval(owner or user), "threat": threat, **extra})
    return services.complete_risk(user, risk)


def _planned_risk(user, plant, threat, owner=None, **extra):
    """Rischio completato con una misura nel piano: si può inviare in approvazione."""
    from apps.risk.models import RiskMitigationPlan

    risk = _completed_risk(user, plant, threat, owner, **extra)
    RiskMitigationPlan.objects.create(assessment=risk, action="EDR", due_date="2030-01-01", owner_external="MSP")
    return risk


def _site_owner(plant, name="owner"):
    return _user(name, scope="single_plant", plants=[plant])


def _close_coverage(user, plant, skip=()):
    for pair in services.register_coverage(plant)["pairs"]:
        if pair["state"] == "missing" and pair["threat_code"] not in skip:
            threat = ThreatCatalogEntry.objects.get(pk=pair["threat_id"])
            services.mark_not_applicable(user, plant, pair["asset_type"], threat, "Non presente nel sito")


# ── valutazione ──────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_create_requires_open_cycle(org_user, plant, threats):
    with pytest.raises(ValidationError):
        services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["malware"]})


@pytest.mark.django_db
def test_create_computes_classes_and_defaults(org_user, plant, threats, cycle):
    risk = services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["malware"]})
    assert risk.status == "bozza"
    assert risk.cycle == cycle and risk.evaluated_in_cycle == cycle
    assert risk.impact == 4 and risk.matrix_class == "critical" and risk.current_class == "critical"
    assert risk.expected_class == "high"
    # nome vuoto: si mostra il titolo della minaccia nella lingua di chi guarda
    assert risk.name == "" and services.risk_label(risk, "it") == "Minaccia IN_MAL"
    assert risk.nis2_in_scope is True


@pytest.mark.django_db
def test_threat_must_apply_to_asset_type(org_user, plant, threats, cycle):
    with pytest.raises(ValidationError):
        services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["fire"]})


@pytest.mark.django_db
def test_confidentiality_floor_from_information_class(org_user, plant, threats, cycle):
    from apps.risk.models import InformationClass

    ic = InformationClass.objects.create(plant=plant, name="Disegni OEM", confidentiality="very_high")
    risk = services.create_risk(org_user, plant, {
        **_eval(org_user, impact_operational=2, probability=2), "threat": threats["malware"],
        "information_classes": [ic],
    })
    assert risk.impact == 5 and risk.current_class == "high"
    # l'override non scende sotto la soglia di riservatezza
    risk = services.update_risk(org_user, risk, {"class_override": -1, "override_rationale": "Al limite"})
    assert risk.current_class == "high"


def _floor_risk(user, plant, threat, **extra):
    """Rischio su informazioni very_high (soglia 5) con impatto atteso 3."""
    from apps.risk.models import InformationClass

    ic = InformationClass.objects.create(plant=plant, name="Disegni OEM", confidentiality="very_high")
    return services.create_risk(user, plant, {
        **_eval(user, impact_operational=3, probability=4, expected_probability=2, expected_impact=3),
        "threat": threat, "information_classes": [ic], **extra,
    })


@pytest.mark.django_db
def test_expected_class_respects_confidentiality_floor(org_user, plant, threats, cycle):
    risk = _floor_risk(org_user, plant, threats["malware"])
    # 2 × 3 sarebbe Medium, ma l'impatto atteso non scende sotto 5: 2 × 5
    assert risk.expected_class == services.risk_class(2, 5) == "high"
    assert risk.expected_impact == 3  # il valore inserito resta: lo segnala la completezza
    with pytest.raises(ValidationError, match="5"):
        services.complete_risk(org_user, risk)
    risk = services.update_risk(org_user, risk, {"expected_impact": 5})
    risk = services.complete_risk(org_user, risk)
    # applicando l'atteso si ottiene proprio la classe attesa disegnata
    plan = RiskMitigationPlan.objects.create(assessment=risk, action="DLP", due_date=datetime.date(2030, 1, 1),
                                             completed_at=timezone.now())
    services.verify_mitigation_plan(org_user, plan, "Test DLP superato")
    expected = risk.expected_class
    risk = services.apply_expected_risk(org_user, risk)
    assert risk.current_class == expected


@pytest.mark.django_db
def test_expected_floor_only_for_confidentiality_threats(org_user, plant, threats, cycle):
    from apps.risk.models import InformationClass

    ic = InformationClass.objects.create(plant=plant, name="Disegni OEM", confidentiality="very_high")
    risk = services.create_risk(org_user, plant, {
        **_eval(org_user, asset_type="SEDE", expected_probability=2, expected_impact=3),
        "threat": threats["fire"], "information_classes": [ic],
    })
    assert risk.expected_class == services.risk_class(2, 3)
    # rischio accettato: l'impatto atteso non conta, la completezza non lo blocca
    risk = _floor_risk(org_user, plant, threats["malware"], treatment="accettare",
                       treatment_rationale="Costi superiori ai benefici")
    services.complete_risk(org_user, risk)


@pytest.mark.django_db
def test_consistency_flags_expected_below_floor(org_user, plant, threats, cycle):
    risk = _floor_risk(org_user, plant, threats["malware"])
    found = [f for f in services.register_consistency_checks(plant) if f["code"] == "expected_below_floor"]
    assert found and found[0]["risk_id"] == str(risk.pk) and found[0]["severity"] == "error"
    assert found[0]["params"] == {"expected": 3, "floor": 5}


@pytest.mark.django_db
def test_ai_expected_impact_raised_to_floor(plant, threats):
    from apps.risk.models import InformationClass

    ic = InformationClass.objects.create(plant=plant, name="Disegni OEM", confidentiality="high")
    out = services.clamp_ai_expected_impact({"expected_impact": 2, "treatment_rationale": "MFA"},
                                            threats["malware"], [ic.pk])
    assert out["expected_impact"] == 4 and out["treatment_rationale"].startswith("MFA ")
    same = {"expected_impact": 2}
    assert services.clamp_ai_expected_impact(same, threats["fire"], [ic.pk]) == same
    assert services.clamp_ai_expected_impact(same, threats["malware"], []) == same


@pytest.mark.django_db
def test_migration_recomputes_stored_expected_class(org_user, plant, threats, cycle):
    import importlib

    from django.apps import apps as django_apps

    from apps.risk.models import RiskAssessment

    risk = _floor_risk(org_user, plant, threats["malware"])
    RiskAssessment.objects.filter(pk=risk.pk).update(expected_class="medium")  # valore pre-correzione
    migration = importlib.import_module("apps.risk.migrations.0021_expected_class_confidentiality_floor")
    migration.forwards(django_apps, None)
    risk.refresh_from_db()
    assert risk.expected_class == "high" and risk.expected_impact == 3


@pytest.mark.django_db
def test_complete_validates_rationales_and_expected(org_user, plant, threats, cycle):
    risk = services.create_risk(org_user, plant, {
        **_eval(org_user, probability_rationale="", expected_probability=None), "threat": threats["malware"],
    })
    with pytest.raises(ValidationError):
        services.complete_risk(org_user, risk)
    risk = services.update_risk(org_user, risk, {"probability_rationale": "Statistiche", "expected_probability": 2})
    risk = services.complete_risk(org_user, risk)
    assert risk.status == "completato" and risk.plan_due_date is not None


@pytest.mark.django_db
def test_accepting_high_needs_cost_benefit(org_user, plant, threats, cycle):
    risk = services.create_risk(org_user, plant, {
        **_eval(org_user, treatment="accettare", expected_probability=None, expected_impact=None),
        "threat": threats["malware"],
    })
    with pytest.raises(ValidationError):
        services.complete_risk(org_user, risk)
    services.update_risk(org_user, risk, {"treatment_rationale": "Costo misure superiore al danno atteso"})
    assert services.complete_risk(org_user, risk).status == "completato"


@pytest.mark.django_db
def test_complete_high_creates_task_and_notification(org_user, plant, threats, cycle,
                                                     django_capture_on_commit_callbacks, monkeypatch):
    from apps.tasks.models import Task

    fired = []
    monkeypatch.setattr("apps.notifications.resolver.fire_notification",
                        lambda event, **kw: fired.append(event))
    with django_capture_on_commit_callbacks(execute=True):
        risk = _completed_risk(org_user, plant, threats["malware"])
    assert Task.objects.filter(source_id=risk.pk, source_module="M06").count() == 1
    assert fired == ["risk_red"]


@pytest.mark.django_db
def test_edit_resets_to_draft_and_requires_cycle(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"])
    risk = services.update_risk(org_user, risk, {"probability": 2})
    assert risk.status == "bozza" and risk.current_class == "high"
    services.complete_risk(org_user, risk)
    cycle.status = "in_approvazione"
    cycle.save()
    with pytest.raises(ValidationError):
        services.update_risk(org_user, risk, {"probability": 3})


@pytest.mark.django_db
def test_legacy_risk_is_read_only(org_user, plant, threats, cycle):
    legacy_cycle = RiskAssessmentCycle.objects.create(plant=plant, kind="legacy", status="archiviato",
                                                      started_at=timezone.now())
    legacy = RiskAssessment.objects.create(plant=plant, cycle=legacy_cycle, name="Vecchio", status="completato")
    with pytest.raises(ValidationError):
        services.update_risk(org_user, legacy, {"probability": 2})
    assert not services.register_queryset(plant).filter(pk=legacy.pk).exists()


@pytest.mark.django_db
def test_site_user_cannot_write_other_site_or_group(site_user, other_plant, threats, org_user):
    services.start_cycle(org_user, other_plant, "primo")
    with pytest.raises(PermissionDenied):
        services.create_risk(site_user, other_plant, {**_eval(site_user), "threat": threats["malware"]})
    with pytest.raises(PermissionDenied):
        services.create_risk(site_user, None, {**_eval(site_user), "threat": threats["malware"]})


# ── copertura e cicli ────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_coverage_and_not_applicable(org_user, plant, threats, cycle):
    cov = services.register_coverage(plant)
    assert cov["asset_types"] == ["IT", "SEDE", "PERSONALE"]
    assert cov["total"] == 3 and cov["missing"] == 3
    _completed_risk(org_user, plant, threats["malware"])
    services.mark_not_applicable(org_user, plant, "SEDE", threats["fire"], "Locali senza materiale combustibile")
    cov = services.register_coverage(plant)
    states = {p["threat_code"]: p["state"] for p in cov["pairs"]}
    assert states == {"IN_MAL": "evaluated", "LO_FUO": "not_applicable", "PE_IPR": "missing"}
    with pytest.raises(ValidationError):
        services.mark_not_applicable(org_user, plant, "PERSONALE", threats["staff"], "  ")


@pytest.mark.django_db
def test_submit_requires_full_coverage_then_approve_freezes_snapshot(org_user, plant, threats, cycle):
    from apps.governance.models import SecurityCommittee

    _planned_risk(org_user, plant, threats["malware"])
    with pytest.raises(ValidationError):
        services.submit_cycle(org_user, cycle)
    _close_coverage(org_user, plant)
    services.submit_cycle(org_user, cycle)
    body = SecurityCommittee.objects.create(name="CdA", committee_type="cda")
    with pytest.raises(ValidationError):
        services.approve_cycle(org_user, cycle, body=None)
    services.approve_cycle(org_user, cycle, body=body)
    cycle.refresh_from_db()
    assert cycle.status == "approvato" and cycle.approved_by_body == body
    assert len(cycle.snapshot["risks"]) == 3
    assert cycle.snapshot["coverage"]["pct"] == 100.0
    assert services.approved_cycle(plant) == cycle


@pytest.mark.django_db
def test_periodic_review_requires_confirmation(org_user, plant, threats, cycle):
    from apps.governance.models import SecurityCommittee

    risk = _planned_risk(org_user, plant, threats["malware"])
    _close_coverage(org_user, plant)
    services.submit_cycle(org_user, cycle)
    services.approve_cycle(org_user, cycle, body=SecurityCommittee.objects.create(name="CdA"))
    review = services.start_cycle(org_user, plant, "periodico")
    errors = services.cycle_submission_errors(review)
    assert any("confermati" in e for e in errors)
    for r in services.register_queryset(plant).filter(applicable=True):
        services.confirm_risk(org_user, r)
    assert services.cycle_submission_errors(review) == []
    services.submit_cycle(org_user, review)
    services.approve_cycle(org_user, review, body=SecurityCommittee.objects.create(name="CdA 2"))
    cycle.refresh_from_db()
    assert cycle.status == "archiviato"
    risk.refresh_from_db()
    assert risk.evaluated_in_cycle == review


@pytest.mark.django_db
def test_site_approval_in_centralised_model_needs_org_scope(org_user, site_user, plant, other_plant, threats):
    from apps.governance.models import SecurityCommittee

    cycle = services.start_cycle(site_user, plant, "primo")
    _planned_risk(site_user, plant, threats["malware"])
    _close_coverage(site_user, plant)
    services.submit_cycle(site_user, cycle)
    body = SecurityCommittee.objects.create(name="CdA")
    assert services.resolve_policy(plant)["preset"] == "centralizzato"
    with pytest.raises(PermissionDenied):
        services.approve_cycle(site_user, cycle, body=body)
    services.approve_cycle(org_user, cycle, body=body)


# ── piani, verifica, rischio atteso ──────────────────────────────────────────

@pytest.mark.django_db
def test_apply_expected_only_after_verified_measures(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"])
    plan = RiskMitigationPlan.objects.create(assessment=risk, action="MFA", due_date=datetime.date(2030, 1, 1))
    with pytest.raises(ValidationError):
        services.apply_expected_risk(org_user, risk)
    with pytest.raises(ValidationError):
        services.verify_mitigation_plan(org_user, plan)
    plan.completed_at = timezone.now()
    plan.save()
    services.verify_mitigation_plan(org_user, plan, "Test MFA superato")
    risk = services.apply_expected_risk(org_user, risk)
    assert risk.current_class == "high" and risk.probability == 2


@pytest.mark.django_db
def test_plan_api_verify_and_uncomplete_resets_verification(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"])
    c = _client(org_user)
    res = c.post("/api/v1/risk/mitigation-plans/", {
        "assessment": str(risk.pk), "action": "EDR", "due_date": "2030-01-01", "owner_external": "MSP",
        "completed_at": timezone.now().isoformat(),
    }, format="json")
    assert res.status_code == 201, res.content
    plan_id = res.json()["id"]
    assert c.post(f"/api/v1/risk/mitigation-plans/{plan_id}/verify/", {"note": "ok"}, format="json").status_code == 200
    res = c.post(f"/api/v1/risk/mitigation-plans/{plan_id}/uncomplete/")
    assert res.status_code == 200 and res.json()["verified_at"] is None



@pytest.mark.django_db
def test_plan_api_edit(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"])
    other = services.create_risk(org_user, plant, {"asset_type": "SEDE", "threat": threats["fire"]})
    c = _client(org_user)
    plan_id = c.post("/api/v1/risk/mitigation-plans/", {
        "assessment": str(risk.pk), "action": "EDR", "due_date": "2030-01-01", "owner_external": "MSP",
    }, format="json").json()["id"]
    url = f"/api/v1/risk/mitigation-plans/{plan_id}/"
    res = c.patch(url, {"action": "EDR su postazioni e server", "due_date": "2030-03-01",
                        "expected_effect": "entrambi"}, format="json")
    assert res.status_code == 200 and res.json()["action"] == "EDR su postazioni e server"
    assert res.json()["due_date"] == "2030-03-01"
    # non si sposta su un altro rischio
    assert c.patch(url, {"assessment": str(other.pk)}, format="json").status_code == 400
    # verificata: contenuti bloccati finché non si riapre
    c.patch(url, {"completed_at": timezone.now().isoformat()}, format="json")
    assert c.post(f"{url}verify/", {"note": "ok"}, format="json").status_code == 200
    assert c.patch(url, {"action": "altro"}, format="json").status_code == 400
    c.post(f"{url}uncomplete/")
    assert c.patch(url, {"action": "altro"}, format="json").status_code == 200

# ── accettazione ─────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_medium_accepted_by_risk_owner_alone(org_user, plant, threats, cycle):
    owner = _user("owner1", scope="single_plant", plants=[plant])
    risk = services.create_risk(org_user, plant, {
        **_eval(owner, probability=3, impact_operational=3, treatment="accettare",
                expected_probability=None, expected_impact=None),
        "threat": threats["malware"],
    })
    risk = services.complete_risk(org_user, risk)
    assert risk.current_class == "medium"
    with pytest.raises(ValidationError):
        services.request_acceptance(org_user, risk, rationale="")
    acc = services.request_acceptance(owner, risk, rationale="Rischio residuo tollerabile")
    assert acc.status == "active" and acc.required_roles == ["risk_owner"]
    assert acc.expires_on <= timezone.localdate() + datetime.timedelta(days=360)


@pytest.mark.django_db
def test_high_needs_plant_manager_and_binding_ciso_opinion(org_user, plant, threats, cycle):
    from apps.governance.models import RoleAssignment

    _plant("RG-Z")  # due siti → preset centralizzato
    owner = _user("owner2", scope="single_plant", plants=[plant])
    pm = _user("pm", role="plant_manager", scope="single_plant", plants=[plant])
    ciso = _user("ciso", role="compliance_officer")
    RoleAssignment.objects.create(user=ciso, role="ciso", scope_type="org", valid_from=datetime.date(2024, 1, 1))
    risk = services.create_risk(org_user, plant, {
        **_eval(owner, probability=3, impact_operational=4, treatment="accettare",
                treatment_rationale="Costi sproporzionati", expected_probability=None, expected_impact=None),
        "threat": threats["malware"],
    })
    risk = services.complete_risk(org_user, risk)
    assert risk.current_class == "high"
    acc = services.request_acceptance(owner, risk, rationale="Analisi costi/benefici allegata")
    assert acc.status == "pending" and acc.upper_opinion == "pending"
    assert set(acc.required_roles) == {"risk_owner", "plant_manager"}
    with pytest.raises(ValidationError):
        services.sign_acceptance(owner, acc)  # ha già firmato
    services.sign_acceptance(pm, acc)
    acc.refresh_from_db()
    assert acc.status == "pending"  # manca il parere
    with pytest.raises(ValidationError):
        services.give_opinion(pm, acc, favorable=True)
    services.give_opinion(ciso, acc, favorable=True, note="Coerente con la policy")
    acc.refresh_from_db()
    assert acc.status == "active"


@pytest.mark.django_db
def test_unfavorable_opinion_rejects(org_user, plant, threats, cycle):
    _plant("RG-Y")
    ciso = _user("ciso2", role="compliance_officer")
    risk = _completed_risk(org_user, plant, threats["malware"], probability=3, treatment="accettare",
                           treatment_rationale="x", expected_probability=None, expected_impact=None)
    acc = services.request_acceptance(org_user, risk, rationale="Motivo")
    with pytest.raises(ValidationError):
        services.give_opinion(ciso, acc, favorable=False, note="")
    services.give_opinion(ciso, acc, favorable=False, note="Serve una misura compensativa")
    acc.refresh_from_db()
    assert acc.status == "rejected"


@pytest.mark.django_db
def test_critical_requires_body_resolution(org_user, plant, threats, cycle):
    from apps.governance.models import SecurityCommittee

    risk = _completed_risk(org_user, plant, threats["malware"], probability=5, impact_operational=5,
                           treatment="accettare", treatment_rationale="Eccezione temporanea",
                           expected_probability=None, expected_impact=None)
    assert risk.current_class == "critical"
    acc = services.request_acceptance(org_user, risk, rationale="Piano di rientro in 6 mesi")
    assert acc.requires_body and acc.status == "pending"
    with pytest.raises(ValidationError):
        services.request_acceptance(org_user, risk, rationale="doppia")
    body = SecurityCommittee.objects.create(name="CdA", committee_type="cda")
    services.record_body_decision(org_user, acc, body=body, resolution_ref="Delibera 12/2026")
    acc.refresh_from_db()
    assert acc.status == "active"
    assert acc.expires_on <= timezone.localdate() + datetime.timedelta(days=180)


@pytest.mark.django_db
def test_legal_violation_is_never_acceptable(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"], probability=2, impact_operational=2,
                           treatment="accettare", legal_or_contract_violation=True,
                           expected_probability=None, expected_impact=None)
    with pytest.raises(ValidationError):
        services.request_acceptance(org_user, risk, rationale="Software senza licenza")


@pytest.mark.django_db
def test_class_change_revokes_acceptance(org_user, plant, threats, cycle):
    owner = _site_owner(plant)
    risk = _completed_risk(org_user, plant, threats["malware"], owner=owner, probability=2, impact_operational=3,
                           treatment="accettare", expected_probability=None, expected_impact=None)
    acc = services.request_acceptance(owner, risk, rationale="ok")
    assert acc.status == "active"
    services.update_risk(org_user, risk, {"probability": 4, "treatment_rationale": "Costi sproporzionati"})
    services.complete_risk(org_user, risk)
    acc.refresh_from_db()
    assert acc.status == "revoked"


@pytest.mark.django_db
def test_expiry_job_warns_then_expires(org_user, plant, threats, cycle):
    from apps.tasks.models import Task

    owner = _site_owner(plant)
    risk = _completed_risk(org_user, plant, threats["malware"], owner=owner, probability=2, impact_operational=3,
                           treatment="accettare", expected_probability=None, expected_impact=None)
    acc = services.request_acceptance(owner, risk, rationale="ok",
                                      expires_on=timezone.localdate() + datetime.timedelta(days=10))
    assert services.expire_acceptances() == {"warned": 1, "expired": 0}
    assert services.expire_acceptances() == {"warned": 0, "expired": 0}
    later = timezone.localdate() + datetime.timedelta(days=11)
    assert services.expire_acceptances(later)["expired"] == 1
    acc.refresh_from_db()
    assert acc.status == "expired"
    assert Task.objects.filter(source_id=risk.pk).count() >= 2


@pytest.mark.django_db
def test_overdue_plans_are_notified_then_escalated(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"], probability=3, impact_operational=3)
    plan = RiskMitigationPlan.objects.create(assessment=risk, action="Patch",
                                             due_date=timezone.localdate() - datetime.timedelta(days=5))
    assert services.escalate_overdue_plans() == {"notified": 1, "escalated": 0}
    assert services.escalate_overdue_plans() == {"notified": 0, "escalated": 0}
    later = timezone.localdate() + datetime.timedelta(days=40)
    assert services.escalate_overdue_plans(later)["escalated"] == 1
    plan.refresh_from_db()
    assert plan.escalation_level == 2


@pytest.mark.django_db
def test_acceptance_api_flow(org_user, plant, threats, cycle):
    owner = _site_owner(plant)
    risk = _completed_risk(org_user, plant, threats["malware"], owner=owner, probability=2, impact_operational=3,
                           treatment="accettare", expected_probability=None, expected_impact=None)
    c = _client(owner)
    res = c.post("/api/v1/risk/acceptances/", {"risk": str(risk.pk), "rationale": "Tollerabile"}, format="json")
    assert res.status_code == 201, res.content
    assert res.json()["status"] == "active"
    acc_id = res.json()["id"]
    res = c.post(f"/api/v1/risk/acceptances/{acc_id}/revoke/", {"reason": "Nuova minaccia"}, format="json")
    assert res.status_code == 200 and res.json()["status"] == "revoked"
    assert RiskAcceptance.objects.get(pk=acc_id).close_reason == "Nuova minaccia"


# ── rischi di gruppo ed ereditati ────────────────────────────────────────────

@pytest.mark.django_db
def test_group_risk_inherited_by_sites(org_user, site_user, plant, other_plant, threats):
    group_cycle = services.start_cycle(org_user, None, "primo")
    risk = services.create_risk(org_user, None, {
        **_eval(org_user), "threat": threats["malware"], "affected_plants": [plant, other_plant],
    })
    assert risk.plant is None and risk.cycle == group_cycle
    c = _client(site_user)
    listed = c.get(f"/api/v1/risk/assessments/?plant={plant.pk}&include_inherited=1").json()["results"]
    assert [x["id"] for x in listed] == [str(risk.pk)] and listed[0]["is_inherited"] is True
    assert c.get(f"/api/v1/risk/assessments/?plant={plant.pk}").json()["results"] == []
    # il sito segnala un impatto locale, il gruppo lo recepisce
    report = services.report_local_impact(site_user, risk, plant, local_impact=5, note="Linea unica JIT")
    with pytest.raises(PermissionDenied):
        services.acknowledge_local_impact(site_user, report)
    services.acknowledge_local_impact(org_user, report)
    report.refresh_from_db()
    assert report.status == "recepita"


@pytest.mark.django_db
def test_site_risk_cannot_list_affected_plants(org_user, plant, other_plant, threats, cycle):
    with pytest.raises(ValidationError):
        services.create_risk(org_user, plant, {
            **_eval(org_user), "threat": threats["malware"], "affected_plants": [other_plant],
        })


# ── API registro ─────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_register_api_crud_and_actions(site_user, plant, other_plant, threats, org_user):
    services.start_cycle(site_user, plant, "primo")
    c = _client(site_user)
    payload = {"plant": str(plant.pk), "threat": str(threats["malware"].pk), "asset_type": "IT",
               "probability": 3, "probability_rationale": "Storico", "impact_operational": 3,
               "impact_rationale": "Disagi", "owner": site_user.pk, "treatment": "accettare",
               "treatment_owner_external": "  MSP Srl  ", "business_objectives": [str(_objective().pk)]}
    res = c.post("/api/v1/risk/assessments/", payload, format="json")
    assert res.status_code == 201, res.content
    rid = res.json()["id"]
    assert res.json()["current_class"] == "medium"
    assert res.json()["treatment_owner_name"] == "MSP Srl"
    res = c.patch(f"/api/v1/risk/assessments/{rid}/", {"treatment_owner": site_user.pk}, format="json")
    assert res.json()["treatment_owner_external"] == ""
    assert c.get(f"/api/v1/risk/assessments/{rid}/completeness/").json() == {"errors": []}
    assert c.post(f"/api/v1/risk/assessments/{rid}/complete/").json()["status"] == "completato"
    req = c.get(f"/api/v1/risk/assessments/{rid}/acceptance-requirements/").json()
    assert req["class"] == "medium" and req["roles"][0] == "risk_owner"
    cov = c.get(f"/api/v1/risk/assessments/coverage/?plant={plant.pk}").json()
    assert cov["closed"] == 1
    matrix = c.get(f"/api/v1/risk/assessments/matrix/?plant={plant.pk}").json()
    assert sum(cell["count"] for cell in matrix) == 1 and len(matrix) == 25
    # altro sito: fuori perimetro
    services.start_cycle(org_user, other_plant, "primo")
    res = c.post("/api/v1/risk/assessments/", {**payload, "plant": str(other_plant.pk)}, format="json")
    assert res.status_code == 403
    assert c.delete(f"/api/v1/risk/assessments/{rid}/").status_code == 204
    assert not RiskAssessment.objects.filter(pk=rid).exists()


@pytest.mark.django_db
def test_legacy_listing(org_user, plant):
    legacy_cycle = RiskAssessmentCycle.objects.create(plant=plant, kind="legacy", status="archiviato",
                                                      started_at=timezone.now())
    RiskAssessment.objects.create(plant=plant, cycle=legacy_cycle, name="Vecchio",
                                  legacy_snapshot={"score": 12})
    c = _client(org_user)
    assert c.get(f"/api/v1/risk/assessments/?plant={plant.pk}").json()["results"] == []
    legacy = c.get(f"/api/v1/risk/assessments/?plant={plant.pk}&legacy=1").json()["results"]
    assert legacy[0]["is_legacy"] is True and legacy[0]["legacy_snapshot"]["score"] == 12


@pytest.mark.django_db
def test_cycle_api_submit_and_approve(org_user, plant, threats, cycle):
    from apps.governance.models import SecurityCommittee

    c = _client(org_user)
    assert c.post(f"/api/v1/risk/cycles/{cycle.pk}/submit/").status_code == 400
    _planned_risk(org_user, plant, threats["malware"])
    _close_coverage(org_user, plant)
    assert c.get(f"/api/v1/risk/cycles/{cycle.pk}/submission-check/").json() == {"errors": []}
    assert c.post(f"/api/v1/risk/cycles/{cycle.pk}/submit/").json()["status"] == "in_approvazione"
    body = SecurityCommittee.objects.create(name="CdA")
    res = c.post(f"/api/v1/risk/cycles/{cycle.pk}/approve/", {"body": str(body.pk)}, format="json")
    assert res.status_code == 200 and res.json()["status"] == "approvato"


# ── export e trigger ─────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_excel_export_sheets(org_user, plant, threats, cycle):
    from openpyxl import load_workbook

    _completed_risk(org_user, plant, threats["malware"])
    wb = load_workbook(io.BytesIO(services.generate_risk_excel(plant)))
    assert wb.sheetnames == ["Registro", "Piano di trattamento", "Accettazioni", "Obiettivi aziendali",
                             "Copertura informazioni", "Copertura", "Criteri"]
    reg = wb["Registro"]
    assert reg.cell(row=2, column=3).value == "IN_MAL"
    assert reg.cell(row=2, column=29).value == "Critical"
    res = _client(org_user).get(f"/api/v1/risk/assessments/export/?plant={plant.pk}")
    assert res.status_code == 200


@pytest.mark.django_db
def test_revaluation_triggers_after_approval(org_user, plant, threats, cycle):
    from apps.governance.models import SecurityCommittee
    from apps.incidents.models import Incident

    _planned_risk(org_user, plant, threats["malware"])
    _close_coverage(org_user, plant)
    services.submit_cycle(org_user, cycle)
    services.approve_cycle(org_user, cycle, body=SecurityCommittee.objects.create(name="CdA"))
    assert services.revaluation_triggers(plant) == []
    Incident.objects.create(plant=plant, title="Ransomware", description="x", detected_at=timezone.now(),
                            severity="critica", status="aperto")
    kinds = [t["kind"] for t in services.revaluation_triggers(plant)]
    assert kinds == ["significant_incidents"]


@pytest.mark.django_db
def test_self_managed_risk_needs_plant_manager(org_user, plant, threats, cycle):
    risk = _completed_risk(org_user, plant, threats["malware"], probability=2, impact_operational=3,
                           treatment="accettare", expected_probability=None, expected_impact=None)
    acc = services.request_acceptance(org_user, risk, rationale="ok")
    assert acc.status == "pending" and "plant_manager" in acc.required_roles
    assert services.acceptance_requirements(risk)["added_for_self_management"] is True


@pytest.mark.django_db
def test_self_managed_rule_only_when_owner_would_sign(org_user, plant, threats, cycle):
    """Se la policy fa accettare un altro ruolo (es. solo il CISO), chi ha
    valutato e tratta il rischio non lo accetta da solo: niente Plant Manager."""
    services.save_governance_policy(org_user, None, {"acceptance_matrix": {
        "high": {"roles": ["ciso"], "scope": "plant", "requires_body": False},
        "medium": {"roles": ["ciso"], "scope": "plant", "requires_body": False},
    }})
    risk = _completed_risk(org_user, plant, threats["malware"], probability=2, impact_operational=3,
                           treatment="accettare", expected_probability=None, expected_impact=None)
    req = services.acceptance_requirements(risk)
    assert req["roles"] == ["ciso"] and req["added_for_self_management"] is False
    # responsabile del trattamento esterno: non è autogestito
    risk.treatment_owner_external = "MSP"
    services.save_governance_policy(org_user, None, {"acceptance_matrix": {
        risk.current_class: {"roles": ["risk_owner"], "scope": "plant", "requires_body": False}}})
    assert services.acceptance_requirements(risk)["roles"] == ["risk_owner"]


@pytest.mark.django_db
def test_policy_upper_opinion_none_for_single_site(org_user, plant, threats, cycle):
    RiskGovernancePolicy.objects.create(plant=None, preset="sito_singolo")
    risk = _completed_risk(org_user, plant, threats["malware"], probability=3, impact_operational=4,
                           treatment="accettare", treatment_rationale="x",
                           expected_probability=None, expected_impact=None)
    owner_is_assessor = services.acceptance_requirements(risk)
    assert "plant_manager" in owner_is_assessor["roles"]
    assert owner_is_assessor["upper_opinion"] == "none"


@pytest.mark.django_db
def test_policy_exposes_org_scope_and_untreated_filter(org_user, site_user, plant, threats, cycle):
    res = _client(site_user).get(f"/api/v1/risk/governance-policies/resolved/?plant={plant.pk}")
    assert res.json()["user_org_scope"] is False
    assert _client(org_user).get("/api/v1/risk/governance-policies/resolved/").json()["user_org_scope"] is True
    high = _completed_risk(org_user, plant, threats["malware"])
    _completed_risk(org_user, plant, threats["fire"], asset_type="SEDE", probability=2, impact_operational=2,
                    treatment="accettare", expected_probability=None, expected_impact=None)
    listed = _client(org_user).get(f"/api/v1/risk/assessments/?plant={plant.pk}&untreated_high=1").json()["results"]
    assert [x["id"] for x in listed] == [str(high.pk)]


@pytest.mark.django_db
def test_cycle_export_and_audit_pack_history(org_user, plant, threats, cycle, tmp_path):
    from openpyxl import load_workbook

    from apps.audit_prep.audit_pack import _collect_risk
    from apps.governance.models import SecurityCommittee

    legacy_cycle = RiskAssessmentCycle.objects.create(plant=plant, kind="legacy", status="archiviato",
                                                      started_at=timezone.now())
    RiskAssessment.objects.create(plant=plant, cycle=legacy_cycle, name="Vecchio",
                                  legacy_snapshot={"score": 12, "inherent_score": 20, "treatment": "mitigare"})
    c = _client(org_user)
    assert c.get(f"/api/v1/risk/cycles/{cycle.pk}/export/").status_code == 400  # niente fotografia
    _planned_risk(org_user, plant, threats["malware"])
    _close_coverage(org_user, plant)
    services.submit_cycle(org_user, cycle)
    services.approve_cycle(org_user, cycle, body=SecurityCommittee.objects.create(name="CdA"))
    res = c.get(f"/api/v1/risk/cycles/{cycle.pk}/export/")
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(b"".join(res.streaming_content) if hasattr(res, "streaming_content") else res.content))
    assert wb.sheetnames == ["Valutazione", "Registro", "Obiettivi aziendali", "Criteri"]
    assert wb["Registro"].max_row == 4  # intestazione + 3 coppie della copertura

    out = _collect_risk(tmp_path, plant)
    assert out["approved_cycle"] is True and out["legacy"] == 1
    files = {p.name for p in (tmp_path / "03_risk").iterdir()}
    assert "valutazione_precedente_metodo_superato.csv" in files
    assert any(f.startswith("valutazione_approvata_") for f in files)


# ── cosa richiede di agire ───────────────────────────────────────────────────

@pytest.mark.django_db
def test_register_attention(org_user, plant, threats, cycle):
    today = datetime.date(2030, 6, 1)
    critical = _completed_risk(org_user, plant, threats["malware"])  # P4 × I4
    high = _completed_risk(org_user, plant, threats["fire"], asset_type="SEDE", impact_operational=3)
    accepted = _completed_risk(org_user, plant, threats["staff"], asset_type="PERSONALE", impact_operational=3)
    RiskAcceptance.objects.create(risk=accepted, risk_class="high", status="active", rationale="Costi",
                                  expires_on=today + datetime.timedelta(days=10))
    RiskMitigationPlan.objects.create(assessment=high, action="Sprinkler", due_date=today - datetime.timedelta(days=1))
    RiskMitigationPlan.objects.create(assessment=high, action="Fatto", due_date=today - datetime.timedelta(days=5),
                                      completed_at=timezone.now())

    data = services.register_attention(plant, today=today)
    assert critical.current_class == "critical" and high.current_class == "high"
    assert data["critical_untreated"]["risk_ids"] == [str(critical.pk)]
    assert data["high_untreated"]["risk_ids"] == [str(high.pk)]
    assert data["acceptances_expiring"]["risk_ids"] == [str(accepted.pk)]
    assert data["overdue_measures"] == {"count": 1, "risk_ids": [str(high.pk)], "measures": 1}
    assert services.register_attention(None, today=today)["critical_untreated"]["count"] == 0


@pytest.mark.django_db
def test_register_attention_endpoint(org_user, plant):
    resp = _client(org_user).get(f"/api/v1/risk/assessments/attention/?plant={plant.pk}")
    assert resp.status_code == 200
    assert set(resp.data) == {"critical_untreated", "high_untreated", "acceptances_expiring", "overdue_measures",
                              "inherited"}


@pytest.mark.django_db
def test_legacy_risk_readable_not_writable(org_user, plant):
    from apps.risk.models import RiskAssessment, RiskAssessmentCycle

    legacy = RiskAssessmentCycle.objects.create(plant=plant, kind="legacy", status="archiviato",
                                                started_at=timezone.now(), closed_at=timezone.now())
    risk = RiskAssessment.objects.create(plant=plant, cycle=legacy, name="Vecchio", legacy_snapshot={"score": 12})
    c = _client(org_user)
    resp = c.get(f"/api/v1/risk/assessments/{risk.pk}/")
    assert resp.status_code == 200 and resp.data["is_legacy"] is True
    assert c.patch(f"/api/v1/risk/assessments/{risk.pk}/", {"name": "X"}, format="json").status_code == 404
    listed = c.get(f"/api/v1/risk/assessments/?plant={plant.pk}").data["results"]
    assert str(risk.pk) not in [r["id"] for r in listed]


@pytest.mark.django_db
def test_display_name_follows_viewer_language(org_user, plant, cycle):
    threat = ThreatCatalogEntry.objects.create(
        code="IN_PHI", asset_types=["IT"], cia=["C"], source="catalog",
        translations={"it": {"title": "Phishing mirato"}, "en": {"title": "Targeted phishing"},
                      "pl": {"title": "Ukierunkowany phishing"}},
    )
    risk = services.create_risk(org_user, plant, {**_eval(org_user), "threat": threat})
    c = _client(org_user)
    url = f"/api/v1/risk/assessments/{risk.pk}/"
    assert c.get(url, HTTP_ACCEPT_LANGUAGE="pl").data["display_name"] == "Ukierunkowany phishing"
    assert c.get(url, HTTP_ACCEPT_LANGUAGE="tr").data["display_name"] == "Targeted phishing"
    services.update_risk(org_user, risk, {"name": "Phishing ufficio acquisti"})
    assert c.get(url, HTTP_ACCEPT_LANGUAGE="pl").data["display_name"] == "Phishing ufficio acquisti"


# ── obiettivi aziendali e informazioni ──────────────────────────────────────

@pytest.mark.django_db
def test_complete_requires_business_objective(org_user, plant, threats, cycle):
    risk = services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["malware"],
                                                  "business_objectives": []})
    # tutto il resto è compilato: l'unica mancanza è l'obiettivo aziendale
    assert len(services.risk_completeness_errors(risk)) == 1
    with pytest.raises(ValidationError):
        services.complete_risk(org_user, risk)
    services.update_risk(org_user, risk, {"business_objectives": [_objective()]})
    assert services.complete_risk(org_user, risk).status == "completato"


@pytest.mark.django_db
def test_site_objective_not_usable_elsewhere(org_user, plant, other_plant, threats, cycle):
    from apps.risk.models import BusinessObjective

    other = BusinessObjective.objects.create(plant=other_plant, name="Altro sito")
    with pytest.raises(ValidationError):
        services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["malware"],
                                               "business_objectives": [other]})


@pytest.mark.django_db
def test_register_objectives_summary(org_user, plant, threats, cycle):
    from apps.risk.models import BusinessObjective

    bo = _objective()
    empty = BusinessObjective.objects.create(name="Reputation", impact_dimensions=["reputational"])
    critical = _completed_risk(org_user, plant, threats["malware"])
    rows = {(r["objective"] or {}).get("name"): r for r in services.register_objectives(plant)}
    assert rows[bo.name]["count"] == 1 and rows[bo.name]["worst_class"] == "critical"
    assert rows[bo.name]["untreated_high"] == 1 and rows[bo.name]["risk_ids"] == [str(critical.pk)]
    assert rows[empty.name]["count"] == 0
    assert rows[None]["count"] == 0
    resp = _client(org_user).get(f"/api/v1/risk/assessments/objectives/?plant={plant.pk}")
    assert resp.status_code == 200 and len(resp.data) == 3


@pytest.mark.django_db
def test_information_coverage(org_user, plant, threats, cycle):
    from apps.risk.models import InformationClass

    secret = InformationClass.objects.create(name="Prototype CAD", confidentiality="very_high")
    InformationClass.objects.create(name="Brochures", confidentiality="low")
    state = {i["name"]: i["state"] for i in services.information_coverage(plant)}
    assert state == {"Prototype CAD": "missing"}
    _completed_risk(org_user, plant, threats["malware"], information_classes=[secret], expected_impact=5)
    item = services.information_coverage(plant)[0]
    assert item["state"] == "evaluated" and item["worst_class"] == "critical"
    # una minaccia senza riservatezza (incendio) non copre la classe
    other = InformationClass.objects.create(name="Customer data", confidentiality="high")
    _completed_risk(org_user, plant, threats["fire"], asset_type="SEDE", information_classes=[other])
    assert {i["name"]: i["state"] for i in services.information_coverage(plant)}["Customer data"] == "missing"


@pytest.mark.django_db
def test_org_wide_supplier_allowed_in_site_risk(org_user, plant, cycle):
    from apps.suppliers.models import Supplier

    threat = ThreatCatalogEntry.objects.create(code="FO_X", asset_types=["FORNITORI"], cia=["A"],
                                               source="catalog", translations={"en": {"title": "Supplier"}})
    supplier = Supplier.objects.create(name="Org supplier")  # nessun sito = di organizzazione
    risk = services.create_risk(org_user, plant, {**_eval(org_user), "asset_type": "FORNITORI",
                                                  "threat": threat, "supplier": supplier})
    assert risk.supplier == supplier


@pytest.mark.django_db
def test_security_objective_links_risks(org_user, plant, other_plant, threats, cycle):
    from apps.governance.models import SecurityObjective

    risk = _completed_risk(org_user, plant, threats["malware"])
    c = _client(org_user)
    base = {"code": "OBJ-R1", "title": "Restore test MES", "origin": "risk_assessment",
            "measure_source": "manual", "unit": "%", "start_date": "2026-01-01", "baseline_value": 0,
            "target_value": 100, "target_direction": "above", "target_date": "2026-12-31",
            "plant": str(plant.pk), "risks": [str(risk.pk)]}
    res = c.post("/api/v1/governance/security-objectives/", base, format="json")
    assert res.status_code == 201, res.content
    assert res.json()["risks_summary"][0]["business_objectives"] == ["OEM supply continuity"]
    detail = c.get(f"/api/v1/risk/assessments/{risk.pk}/").json()
    assert detail["security_objectives_summary"][0]["code"] == "OBJ-R1"
    # un rischio di un altro sito non si collega
    services.start_cycle(org_user, other_plant, "primo")
    foreign = _completed_risk(org_user, other_plant, threats["malware"])
    res = c.post("/api/v1/governance/security-objectives/",
                 {**base, "code": "OBJ-R2", "risks": [str(foreign.pk)]}, format="json")
    assert res.status_code == 400
    assert SecurityObjective.objects.filter(code="OBJ-R2").count() == 0


@pytest.mark.django_db
def test_site_coverage_counts_inherited_group_risks(org_user, plant, other_plant, threats):
    """Un rischio di gruppo ereditato copre la coppia nel sito: niente duplicati."""
    services.start_cycle(org_user, None, "primo")
    group = services.create_risk(org_user, None, {**_eval(org_user), "threat": threats["malware"],
                                                  "affected_plants": [plant]})
    pair = next(p for p in services.register_coverage(plant)["pairs"]
                if p["threat_code"] == "IN_MAL" and p["asset_type"] == "IT")
    assert pair["state"] == "inherited" and pair["risk_ids"] == [str(group.pk)]
    # un rischio proprio del sito sulla stessa minaccia prevale
    services.start_cycle(org_user, plant, "primo")
    own = services.create_risk(org_user, plant, {**_eval(org_user), "threat": threats["malware"]})
    pair = next(p for p in services.register_coverage(plant)["pairs"]
                if p["threat_code"] == "IN_MAL" and p["asset_type"] == "IT")
    assert pair["state"] == "draft" and pair["risk_ids"] == [str(own.pk)]


@pytest.mark.django_db
def test_site_coverage_ignores_group_risks_not_inherited(org_user, plant, other_plant, threats):
    services.start_cycle(org_user, None, "primo")
    services.create_risk(org_user, None, {**_eval(org_user), "threat": threats["malware"],
                                          "affected_plants": [other_plant]})
    pair = next(p for p in services.register_coverage(plant)["pairs"] if p["threat_code"] == "IN_MAL")
    assert pair["state"] == "missing"



@pytest.mark.django_db
def test_inherited_group_risks_shown_apart_not_summed(org_user, plant, other_plant, threats):
    """Regola unica: i rischi di gruppo che riguardano il sito si mostrano a parte
    e non entrano nei numeri del sito; nell'organizzazione contano una volta."""
    services.start_cycle(org_user, None, "primo")
    group = services.complete_risk(org_user, services.create_risk(org_user, None, {
        **_eval(org_user), "threat": threats["malware"], "affected_plants": [plant, other_plant]}))
    services.start_cycle(org_user, plant, "primo")
    _completed_risk(org_user, plant, threats["fire"], asset_type="SEDE")
    assert services.inherited_summary(plant) == {"count": 1, "untreated_high": 1, "risk_ids": [str(group.pk)]}
    row = next(r for r in services.register_objectives(plant) if r["objective"])
    assert row["count"] == 1 and row["inherited_count"] == 1 and row["inherited_risk_ids"] == [str(group.pk)]
    attention = services.register_attention(plant)
    assert str(group.pk) not in attention["critical_untreated"]["risk_ids"]
    assert attention["inherited"]["untreated_high"] == 1
    assert services.register_attention(None)["inherited"] is None



@pytest.mark.django_db
def test_submit_blocked_for_high_risk_without_plan(org_user, plant, threats, cycle):
    """L'organo approva valutazione e piano (ISO 27001 §6.1.3): un High/Critical
    da trattare senza misure non va in approvazione; una misura eliminata non conta."""
    from apps.risk.models import RiskMitigationPlan

    risk = _completed_risk(org_user, plant, threats["malware"])  # Critical da mitigare
    _close_coverage(org_user, plant)
    assert any("piano di trattamento" in e for e in services.cycle_submission_errors(cycle))
    plan = RiskMitigationPlan.objects.create(assessment=risk, action="EDR", due_date="2030-01-01",
                                             owner_external="MSP")
    assert services.cycle_submission_errors(cycle) == []
    plan.soft_delete()
    with pytest.raises(ValidationError):
        services.submit_cycle(org_user, cycle)
