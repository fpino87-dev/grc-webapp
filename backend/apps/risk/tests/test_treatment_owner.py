"""
Responsabile del trattamento (scenario) e responsabile della singola azione di
mitigazione: campo misto utente del portale OPPURE testo libero.
"""
import io

import pytest
from django.contrib.auth import get_user_model
from rest_framework.test import APIClient

User = get_user_model()


@pytest.fixture
def super_admin(db):
    from apps.auth_grc.models import GrcRole, UserPlantAccess
    user = User.objects.create_user(username="sa-to@test.com", email="sa-to@test.com", password="x")
    UserPlantAccess.objects.create(user=user, role=GrcRole.SUPER_ADMIN, scope_type="org")
    return user


@pytest.fixture
def client(super_admin):
    c = APIClient()
    c.force_authenticate(user=super_admin)
    return c


@pytest.fixture
def tech_user(db):
    return User.objects.create_user(
        username="it.tech", email="it.tech@test.com", password="x",
        first_name="Mario", last_name="Rossi",
    )


@pytest.fixture
def plant(db):
    from apps.plants.models import Plant
    return Plant.objects.create(
        code="TO-P", name="Plant TO", country="IT",
        nis2_scope="non_soggetto", status="attivo",
    )


@pytest.fixture
def assessment(db, plant, super_admin):
    from apps.risk.models import RiskAssessment
    return RiskAssessment.objects.create(
        plant=plant, name="Rischio TO", assessment_type="IT",
        probability=3, impact=4, created_by=super_admin,
    )


URL = "/api/v1/risk/assessments/"
PLAN_URL = "/api/v1/risk/mitigation-plans/"


@pytest.mark.django_db
def test_create_with_portal_user(client, plant, super_admin, tech_user):
    res = client.post(URL, {
        "plant": str(plant.pk), "name": "R1", "assessment_type": "IT",
        "probability": 2, "impact": 3,
        "owner": super_admin.pk, "treatment_owner": tech_user.pk,
    }, format="json")
    assert res.status_code == 201, res.data
    assert res.data["treatment_owner"] == tech_user.pk
    assert res.data["treatment_owner_external"] == ""
    assert res.data["treatment_owner_name"] == "Mario Rossi"


@pytest.mark.django_db
def test_create_with_external_text(client, plant):
    res = client.post(URL, {
        "plant": str(plant.pk), "name": "R2", "assessment_type": "OT",
        "probability": 2, "impact": 3,
        "treatment_owner_external": "  MSP Rete Srl  ",
    }, format="json")
    assert res.status_code == 201, res.data
    assert res.data["treatment_owner"] is None
    assert res.data["treatment_owner_external"] == "MSP Rete Srl"
    assert res.data["treatment_owner_name"] == "MSP Rete Srl"


@pytest.mark.django_db
def test_owner_and_treatment_owner_can_coincide(client, assessment, tech_user):
    res = client.patch(f"{URL}{assessment.pk}/", {
        "owner": tech_user.pk, "treatment_owner": tech_user.pk,
    }, format="json")
    assert res.status_code == 200, res.data
    assert res.data["owner"] == res.data["treatment_owner"] == tech_user.pk


@pytest.mark.django_db
def test_user_wins_over_text(client, assessment, tech_user):
    res = client.patch(f"{URL}{assessment.pk}/", {
        "treatment_owner": tech_user.pk, "treatment_owner_external": "Altro",
    }, format="json")
    assert res.status_code == 200
    assert res.data["treatment_owner_external"] == ""


@pytest.mark.django_db
def test_switch_user_to_text_and_partial_patch(client, assessment, tech_user):
    assessment.treatment_owner = tech_user
    assessment.save()
    # Solo testo: sostituisce l'utente.
    res = client.patch(f"{URL}{assessment.pk}/", {"treatment_owner_external": "Integratore OT"}, format="json")
    assert res.data["treatment_owner"] is None
    assert res.data["treatment_owner_name"] == "Integratore OT"
    # PATCH di altri campi: il responsabile non cambia.
    res = client.patch(f"{URL}{assessment.pk}/", {"name": "Rinominato"}, format="json")
    assert res.data["treatment_owner_name"] == "Integratore OT"


@pytest.mark.django_db
def test_mitigation_plan_owner_mixed(client, assessment, tech_user):
    res = client.post(PLAN_URL, {
        "assessment": str(assessment.pk), "action": "Segmentare la rete OT",
        "due_date": "2026-12-31", "owner_external": "Fornitore PLC",
    }, format="json")
    assert res.status_code == 201, res.data
    assert res.data["owner"] is None
    assert res.data["owner_name"] == "Fornitore PLC"
    res = client.patch(f"{PLAN_URL}{res.data['id']}/", {"owner": tech_user.pk}, format="json")
    assert res.data["owner_external"] == ""
    assert res.data["owner_name"] == "Mario Rossi"


@pytest.mark.django_db
def test_excel_export_has_treatment_owner(assessment):
    from openpyxl import load_workbook
    from apps.risk.models import RiskMitigationPlan
    from apps.risk.services import generate_risk_excel

    assessment.treatment_owner_external = "MSP Rete Srl"
    assessment.save()
    RiskMitigationPlan.objects.create(
        assessment=assessment, action="Patch firewall", due_date="2026-12-31",
        owner_external="Fornitore FW",
    )
    ws = load_workbook(io.BytesIO(generate_risk_excel(include_draft=True))).active
    headers = [c.value for c in ws[1]]
    row = {h: c.value for h, c in zip(headers, ws[2], strict=True)}
    assert row["Responsabile trattamento"] == "MSP Rete Srl"
    assert "[Fornitore FW]" in row["Azioni di mitigazione"]
